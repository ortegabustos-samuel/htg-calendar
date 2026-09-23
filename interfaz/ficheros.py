"""
ficheros.py — Lo que la interfaz hace con el disco: escenarios, config.toml y CSV.

Un ESCENARIO es una carpeta `escenarios/<nombre>/` con `input/` (config.toml + los 6 CSV, el mismo
contrato que data/input) y `output/`. El pipeline la lee a través de las variables de entorno
HT_DATOS y HT_SALIDA.
"""
from __future__ import annotations

import csv
import io
import math
import re
import shutil
import tomllib
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
ESCENARIOS = RAIZ / "escenarios"
DATA_ORIGINAL = RAIZ / "data" / "input"

# Orden en el que se muestran; la descripción es para quien no conoce el contrato.
CSVS = {
    "turnos.csv": "Líneas del Plan Funcional: días que operan, horario, horas y demanda",
    "patrones.csv": "Matrices de rotación semanal: una fila por semana",
    # Después de turnos y patrones: sus columnas patron y linea se eligen de esos dos ficheros.
    "trabajadores.csv": "Plantilla: tipo, patrón o línea, vacaciones y fila inicial",
    "capacidades.csv": "Qué líneas puede hacer cada trabajador fuera de lo que se deriva solo",
    "festivos.csv": "Festivos del año y su ámbito (nacional o calendario local)",
    "calendarios_municipio.csv": "Qué calendario de festivos sigue cada municipio",
}

# Parámetros de config.toml: sección, valor por defecto, límites y ayuda. Los límites son los mismos
# que comprueba validar_datos.revisar_config: el formulario no deja escribir algo fuera de rango.
PARAMETROS = {
    "anio": ("horizonte", 2027, 2000, 2100,
             "Año a resolver: el cuadrante va del 1 de enero al 31 de diciembre."),
    "horas_objetivo": ("jornada", 1776, 1, 3000,
                       "Jornada anual del convenio. El exceso de cada trabajador se devuelve en días libres."),
    "descanso_minimo": ("convenio", 12, 0, 24,
                        "Horas mínimas de descanso entre el fin de un turno y el inicio del siguiente (C4)."),
    "horas_max_semana": ("convenio", 48, 1, 168,
                         "Máximo de horas trabajadas en una semana ISO, de lunes a domingo (C6)."),
    "dias_max_semana": ("convenio", 6, 1, 7,
                        "Máximo de días trabajados en una semana ISO (C5)."),
}


# --------------------------------------------------------------------------- #
#  Escenarios
# --------------------------------------------------------------------------- #
def listar_escenarios() -> list[str]:
    if not ESCENARIOS.is_dir():
        return []
    return sorted(p.name for p in ESCENARIOS.iterdir() if (p / "input").is_dir())


def carpeta_entrada(escenario: str) -> Path:
    return ESCENARIOS / escenario / "input"


def carpeta_salida(escenario: str) -> Path:
    return ESCENARIOS / escenario / "output"


def nombre_valido(nombre: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9_\-]{1,40}", nombre))


def crear_escenario(nombre: str, copiar_de: Path | None) -> None:
    """Crea la carpeta del escenario. Con `copiar_de` arranca con sus ficheros (config.toml y los
    CSV que haya); sin él, vacía."""
    entrada = carpeta_entrada(nombre)
    if entrada.exists():
        raise FileExistsError(f"ya existe un escenario llamado {nombre!r}")
    entrada.mkdir(parents=True)
    carpeta_salida(nombre).mkdir()
    if copiar_de is not None:
        for f in ["config.toml", *CSVS]:
            if (copiar_de / f).is_file():
                shutil.copy2(copiar_de / f, entrada / f)


# --------------------------------------------------------------------------- #
#  config.toml
# --------------------------------------------------------------------------- #
def config_por_defecto() -> dict:
    return {clave: p[1] for clave, p in PARAMETROS.items()} | {"grupos_rigidos": []}


def leer_config(contenido: bytes) -> dict:
    """Aplana las secciones de config.toml a {parámetro: valor}. Lanza ValueError si no es TOML."""
    try:
        datos = tomllib.loads(contenido.decode("utf-8-sig"))
    except (tomllib.TOMLDecodeError, UnicodeDecodeError) as e:
        raise ValueError(f"no es un TOML válido: {e}") from e
    plano: dict = {}
    for seccion in datos.values():
        if isinstance(seccion, dict):
            plano.update(seccion)
    return plano


def config_de_escenario(escenario: str) -> dict:
    """Lo que haya en config.toml sobre los valores por defecto (así un fichero incompleto no deja
    el formulario a medias)."""
    ruta = carpeta_entrada(escenario) / "config.toml"
    cfg = config_por_defecto()
    if ruta.is_file():
        try:
            cfg |= leer_config(ruta.read_bytes())
        except ValueError:
            pass                                # ilegible: el formulario arranca con los valores por defecto
    return cfg


def escribir_config(escenario: str, cfg: dict) -> str:
    """Genera config.toml con las mismas secciones y comentarios que el de data/input."""
    grupos = ", ".join(f'"{g}"' for g in cfg["grupos_rigidos"])
    texto = f"""# Generado desde la interfaz (interfaz/app.py)
[horizonte]
# Año a resolver: el cuadrante va del 1 de enero al 31 de diciembre.
anio = {int(cfg["anio"])}

[jornada]
# Jornada anual objetivo (h).
horas_objetivo = {int(cfg["horas_objetivo"])}

[convenio]
descanso_minimo = {int(cfg["descanso_minimo"])}       # descanso mínimo entre jornadas (h)      — C4
horas_max_semana = {int(cfg["horas_max_semana"])}      # máx. horas por semana ISO              — C6
dias_max_semana = {int(cfg["dias_max_semana"])}        # máx. días trabajados por semana ISO    — C5

[libranzas]
# Grupos cuyo descanso no se fracciona: se cede el ciclo entero (bloque de trabajo + descanso).
grupos_rigidos = [{grupos}]
"""
    (carpeta_entrada(escenario) / "config.toml").write_text(texto, encoding="utf-8")
    return texto


def patrones_declarados(escenario: str) -> list[str]:
    """Ids de patrón de patrones.csv, para ofrecerlos como grupos rígidos."""
    ruta = carpeta_entrada(escenario) / "patrones.csv"
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return sorted({f.get("patron", "") for f in csv.DictReader(fh)} - {"", None})


# --------------------------------------------------------------------------- #
#  CSV
# --------------------------------------------------------------------------- #
def normalizar_csv(contenido: bytes) -> tuple[bytes, list[str]]:
    """Arregla solo la CODIFICACIÓN de un CSV subido, nunca su contenido, y dice qué ha cambiado.

    Son los tres tropiezos de guardar un CSV desde Excel en español, y ninguno lo detectaría el
    validador: el cargador lee en utf-8 sin BOM (con BOM la primera columna pasa a llamarse
    '\\ufeffid_turno' y el fichero entero se lee mal), y un separador ';' o una codificación
    Windows lo dejan ilegible. Los espacios sobrantes en celdas NO se tocan: el validador los
    señala y el arreglo va en el fichero, que es la fuente de verdad."""
    cambios: list[str] = []
    if contenido.startswith(b"\xef\xbb\xbf"):
        contenido = contenido[3:]
        cambios.append("se ha quitado la marca BOM que añade Excel")
    try:
        texto = contenido.decode("utf-8")
    except UnicodeDecodeError:
        texto = contenido.decode("cp1252")
        cambios.append("estaba en codificación Windows (cp1252) y se ha pasado a UTF-8")
    cabecera = texto.splitlines()[0] if texto else ""
    if ";" in cabecera and "," not in cabecera:
        filas = list(csv.reader(io.StringIO(texto), delimiter=";"))
        salida = io.StringIO()
        csv.writer(salida, lineterminator="\n").writerows(filas)
        texto = salida.getvalue()
        cambios.append("estaba separado por punto y coma y se ha pasado a comas "
                       "(revisa que los decimales usen punto: 8.0, no 8,0)")
    return texto.encode("utf-8"), cambios


# --------------------------------------------------------------------------- #
#  Excel (.xlsx)
# --------------------------------------------------------------------------- #
# Un .xlsx se acepta con la MISMA estructura que el CSV (fila 1 = cabecera, una fila por registro)
# y se convierte a CSV al subirlo: el pipeline solo sabe leer CSV, y así el escenario guarda un
# único formato. Un libro puede traer varios ficheros a la vez, una hoja por fichero, con la hoja
# llamada como el CSV sin extensión (turnos, patrones, trabajadores…).
def es_xlsx(contenido: bytes) -> bool:
    return contenido[:4] == b"PK\x03\x04"               # un .xlsx es un zip


def _celda(valor) -> str:
    """Una celda de Excel al texto que espera el cargador. Excel no guarda texto sino tipos, y los
    tres que romperían la carga son: fechas (el cargador quiere DD/MM/AAAA), horas (H:MM) y
    enteros, que openpyxl devuelve como float (int("1.0") falla en lv, dem, fila…)."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "1" if valor else "0"
    if isinstance(valor, datetime):
        # Una hora sola en Excel es una fecha del 30/12/1899 (el día cero); cualquier otra es fecha.
        if valor.date() in (date(1899, 12, 30), date(1899, 12, 31), date(1900, 1, 1)):
            return f"{valor.hour}:{valor.minute:02d}"
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, time):
        return f"{valor.hour}:{valor.minute:02d}"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _hoja_a_csv(hoja) -> bytes:
    filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    if not filas:
        raise ValueError(f"la hoja '{hoja.title}' está vacía")
    cabecera = filas[0]
    # Excel suele arrastrar columnas y filas vacías al final: se corta donde acaba la cabecera.
    ancho = max((i + 1 for i, c in enumerate(cabecera) if c not in (None, "")), default=0)
    if not ancho:
        raise ValueError(f"la hoja '{hoja.title}' no tiene cabecera en la fila 1")
    salida = io.StringIO()
    escritor = csv.writer(salida, lineterminator="\n")
    for fila in filas:
        celdas = [_celda(v) for v in (fila + [None] * ancho)[:ancho]]
        if any(celdas):
            escritor.writerow(celdas)
    return salida.getvalue().encode("utf-8")


def _abrir_xlsx(contenido: bytes):
    try:
        # data_only: el valor calculado de las fórmulas, no la fórmula.
        return load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as e:
        raise ValueError(f"no es un Excel .xlsx válido: {e}") from e


def xlsx_a_csv(contenido: bytes, nombre: str) -> tuple[bytes, str]:
    """El CSV `nombre` sacado de un .xlsx: la hoja que se llame como él, o si no la primera.
    Devuelve (csv, hoja usada)."""
    libro = _abrir_xlsx(contenido)
    hojas = {h.lower(): h for h in libro.sheetnames}
    hoja = libro[hojas.get(Path(nombre).stem.lower(), libro.sheetnames[0])]
    return _hoja_a_csv(hoja), hoja.title


def csvs_de_xlsx(contenido: bytes) -> dict[str, bytes]:
    """Todos los CSV que trae un libro: cada hoja llamada como uno de ellos (sin extensión)."""
    libro = _abrir_xlsx(contenido)
    por_nombre = {Path(n).stem.lower(): n for n in CSVS}
    return {por_nombre[h.lower()]: _hoja_a_csv(libro[h])
            for h in libro.sheetnames if h.lower() in por_nombre}


def guardar_fichero(escenario: str, nombre: str, contenido: bytes) -> list[str]:
    """Guarda un fichero subido en el escenario. Devuelve los cambios de formato aplicados.
    Lanza ValueError si un config.toml no es TOML o un .xlsx no se puede leer."""
    cambios: list[str] = []
    if nombre == "config.toml":
        leer_config(contenido)
    elif es_xlsx(contenido):
        contenido, hoja = xlsx_a_csv(contenido, nombre)
        cambios.append(f"convertido desde Excel (hoja '{hoja}')")
    else:
        contenido, cambios = normalizar_csv(contenido)
    (carpeta_entrada(escenario) / nombre).write_bytes(contenido)
    return cambios


# --------------------------------------------------------------------------- #
#  Edición manual: patrones, trabajadores, festivos y calendarios_municipio
# --------------------------------------------------------------------------- #
# Los CSV que se pueden teclear en la interfaz en vez de exigir un fichero hecho a mano fuera
# de ella. Se guardan en el mismo escenario, así que quedan versionados junto al resto para
# reproducibilidad.
DIAS = ["Lunes", "Martes", "Miercoles", "Jueves", "Viernes", "Sabado", "Domingo"]      # mismo orden que cargar_datos.DIAS
DESCANSOS = ["DO", "LIBRE"]                                   # mismos valores que cargar_datos.DESCANSOS

EDITABLES = {
    "festivos.csv": ["fecha", "ambito"],
    "calendarios_municipio.csv": ["municipio", "calendario_festivos"],
    "patrones.csv": ["patron", "fila", *DIAS],
    "trabajadores.csv": ["id_trab", "nombre", "municipio", "tipo", "patron", "vac1_inicio",
                         "vac2_inicio", "linea", "fila_inicial", "factor_jornada"],
    "turnos.csv": ["id_turno", "municipio", "lv", "sabado", "domingo", "festivo", "hora_entrada",
                   "hora_salida", "horas_computadas", "dem"],
    "capacidades.csv": ["id_trab", "id_turno", "lv", "sab", "dom", "fest", "v"],
}
TIPOS_TRABAJADOR = ["fijo", "patron", "correturno"]

# Columnas que el editor maneja con un tipo propio en vez de texto libre: fecha (DD/MM/AAAA en el
# CSV), hora (texto validado H:MM de 24 h), entero, decimal o casilla (0/1 en el CSV).
FECHAS = {"festivos.csv": ["fecha"], "trabajadores.csv": ["vac1_inicio", "vac2_inicio"]}
HORAS = {"turnos.csv": ["hora_entrada", "hora_salida"]}
ENTEROS = {"patrones.csv": ["fila"], "trabajadores.csv": ["fila_inicial"], "turnos.csv": ["dem"],
           "capacidades.csv": ["v"]}
DECIMALES = {"turnos.csv": ["horas_computadas"]}
CASILLAS = {"turnos.csv": ["lv", "sabado", "domingo", "festivo"],
            "capacidades.csv": ["lv", "sab", "dom", "fest"]}


# factor_jornada: en el CSV es la fracción (0, 1] que lee cargar_datos, vacía = jornada completa;
# en la interfaz es un porcentaje, que es como viene en un contrato.
def porcentaje_jornada(texto: str) -> float:
    """CSV -> porcentaje. Un valor > 1 se toma como un porcentaje tecleado a mano en el CSV (75 en
    vez de 0.75): el cargador lo rechazaría, y aquí se entiende sin ambigüedad. Vacío o ilegible
    -> 100: es lo que el cargador ya hace con una celda vacía."""
    try:
        valor = float(str(texto).strip().replace(",", "."))
    except ValueError:
        return 100.0
    if valor <= 0 or valor > 100:
        return 100.0
    return valor if valor > 1 else valor * 100


def fraccion_jornada(porcentaje: float | None) -> str:
    """Porcentaje -> CSV. 100 % o una celda que se ha dejado vacía en la tabla (llega como None o
    NaN) se escriben como celda vacía, que el cargador lee como jornada completa."""
    if porcentaje is None or math.isnan(porcentaje) or porcentaje >= 100:
        return ""
    return f"{porcentaje / 100:.4f}".rstrip("0").rstrip(".")


def filas_editables(escenario: str, nombre: str) -> list[dict]:
    """Lee un CSV editable a mano como lista de filas. Vacío si el escenario aún no lo tiene."""
    columnas = EDITABLES[nombre]
    ruta = carpeta_entrada(escenario) / nombre
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return [{c: fila.get(c, "") for c in columnas} for fila in csv.DictReader(fh)]


def trabajadores_declarados(escenario: str) -> list[str]:
    """Ids de trabajadores.csv: los únicos que pueden aparecer en capacidades.csv."""
    ruta = carpeta_entrada(escenario) / "trabajadores.csv"
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return sorted({(fila.get("id_trab") or "").strip() for fila in csv.DictReader(fh)} - {""})


def turnos_declarados(escenario: str) -> list[str]:
    """Ids de turno de turnos.csv. Es la lista contra la que se restringen las celdas de
    lun..dom al editar patrones.csv a mano: solo un turno real, o DO, o LIBRE."""
    ruta = carpeta_entrada(escenario) / "turnos.csv"
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return sorted({(fila.get("id_turno") or "").strip() for fila in csv.DictReader(fh)} - {""})


def municipios_validos(escenario: str) -> list[str]:
    """Municipios que aparecen en turnos.csv: los que de verdad necesitan un calendario. Es la
    lista contra la que se restringe la columna 'municipio' al editar a mano. Si turnos.csv no
    está cargado en este escenario, no hay lista y se deja escribir libremente."""
    ruta = carpeta_entrada(escenario) / "turnos.csv"
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return sorted({(fila.get("municipio") or "").strip() for fila in csv.DictReader(fh)} - {""})


def ambitos_declarados(escenario: str) -> list[str]:
    """Ámbitos ya tecleados en festivos.csv. calendario_festivos va emparejado con 'ambito': solo
    tiene efecto si coincide letra por letra con uno de estos, así que es la única lista válida
    para restringir esa columna al editar calendarios_municipio.csv a mano."""
    ruta = carpeta_entrada(escenario) / "festivos.csv"
    if not ruta.is_file():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return sorted({(fila.get("ambito") or "").strip() for fila in csv.DictReader(fh)} - {""})


def filas_limpias(nombre: str, filas: list[dict]) -> list[dict]:
    """Las filas tal como se escribirán en el CSV. Se descartan las que aún no tienen su primera
    columna (el id): con el guardado automático, una fila recién añadida se guarda en cuanto se le
    pone el id, no antes — si no, sus casillas a 0 la harían parecer rellena y se escribiría una
    fila sin id."""
    columnas = EDITABLES[nombre]
    limpias = [{c: str(fila.get(c, "") or "").strip() for c in columnas} for fila in filas]
    return [f for f in limpias if f[columnas[0]]]


def guardar_filas(escenario: str, nombre: str, filas: list[dict]) -> None:
    """Escribe las filas tecleadas a mano como CSV, en el mismo formato que espera el cargador."""
    salida = io.StringIO()
    escritor = csv.DictWriter(salida, fieldnames=EDITABLES[nombre], lineterminator="\n")
    escritor.writeheader()
    escritor.writerows(filas_limpias(nombre, filas))
    (carpeta_entrada(escenario) / nombre).write_text(salida.getvalue(), encoding="utf-8")
